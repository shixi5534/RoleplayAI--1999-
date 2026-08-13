"""多格式导入管线 + 角色提示词系统 测试。

覆盖：预处理 / Markdown 标题感知分块 / 文件-JSON-CSV-SQLite-PDF-DOCX 导入 /
     知识范围解析 / 混合检索（稠密 + BM25）重排 / RAG 上下文注入 /
     角色卡新字段（tone/knowledge_scope）。
"""
import sqlite3
from pathlib import Path

import pytest

from roleplay.core.knowledge import (
    CharacterStore,
    HashingEmbedder,
    KnowledgeBase,
    chunk_markdown,
    ingest_file,
    ingest_sqlite,
    lore_namespace,
    preprocess_text,
)
from roleplay.core.knowledge.vector_store import _BM25
from roleplay.core.persona_prompt import (
    DEFAULT_RETRIEVAL_NAMESPACES,
    build_rag_context,
    build_roleplay_prompt,
    resolve_knowledge_namespaces,
)
from roleplay.models.character import CharacterCard
from roleplay.core.rag.base import RetrievedChunk


def _write_minimal_pdf(path: Path, text: str) -> None:
    """生成一个极简、合法的单页 PDF（标准 Helvetica，仅 ASCII），用于验证加载器管线。

    不依赖第三方写入器，规避 pypdf/reportlab 的文本写入 API 差异与 CJK 字体依赖。
    """
    esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({esc}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0] * (len(objs) + 1)
    for i, body in enumerate(objs, start=1):
        offsets[i] = len(out)
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
    out += f"startxref\n{xref_offset}\n%%EOF\n".encode()
    path.write_bytes(out)


@pytest.fixture()
def kb(tmp_path):
    return KnowledgeBase(HashingEmbedder(dim=64), persist_dir=tmp_path / "kb")


# ── 预处理 ──
def test_preprocess_normalizes_whitespace_and_controls():
    raw = "第一行\r\n\r\n\r\n  第二行\t带  制表\u0000符 "
    out = preprocess_text(raw)
    assert "\r" not in out and "\u0000" not in out
    assert "\n\n\n" not in out  # 连续空行压缩
    assert out.startswith("第一行")


# ── Markdown 标题感知分块 ──
def test_chunk_markdown_carries_section_path():
    md = "# 角色\n\n## 性格\n她很平静。\n\n## 背景\n她被绑架过。"
    chunks = chunk_markdown(md)
    sections = [meta["section"] for _, meta in chunks]
    assert any("角色 > 性格" in s for s in sections)
    assert any("角色 > 背景" in s for s in sections)
    # 节标题拼进块首，增强召回
    assert any(text.startswith("【角色 > 性格】") for text, _ in chunks)


# ── 文件导入 ──
def test_ingest_markdown_file_retrievable(kb, tmp_path):
    f = tmp_path / "lore.md"
    f.write_text("# 设定\n\n## 发条装置\n发条装置令大脑自动生成人格。", encoding="utf-8")
    r = ingest_file(kb, f, namespace="lore_test", character_id="test")
    assert r["stored"] >= 1 and r["format"] == "md"
    hits = kb.search("发条装置", top_k=2, namespaces=["lore_test"])
    assert hits and "发条装置" in hits[0].text
    assert hits[0].metadata.get("character_id") == "test"
    assert "发条装置" in hits[0].metadata.get("section", "")


def test_ingest_json_and_csv(kb, tmp_path):
    j = tmp_path / "info.json"
    j.write_text('{"角色": {"代号": "飞蛾"}}', encoding="utf-8")
    c = tmp_path / "rows.csv"
    c.write_text("name,alias\n无名者,格蕾丝\n", encoding="utf-8")
    assert ingest_file(kb, j, namespace="docs")["stored"] >= 1
    assert ingest_file(kb, c, namespace="docs")["stored"] >= 1
    hits = kb.search("飞蛾", top_k=3, namespaces=["docs"])
    assert any("飞蛾" in h.text for h in hits)


def test_ingest_unsupported_format_raises(kb, tmp_path):
    f = tmp_path / "x.exe"
    f.write_text("bin", encoding="utf-8")
    with pytest.raises(ValueError):
        ingest_file(kb, f)


# ── SQLite 导入 ──
def test_ingest_sqlite_rows(kb, tmp_path):
    db = tmp_path / "lore.db"
    conn = sqlite3.connect(db)
    conn.execute("CREATE TABLE events (title TEXT, detail TEXT)")
    conn.execute("INSERT INTO events VALUES ('77号往事', '凯拉在蓝手帕旅馆被误认')")
    conn.commit()
    conn.close()
    r = ingest_sqlite(kb, db, table="events", namespace="lore_test")
    assert r["stored"] == 1
    hits = kb.search("凯拉 蓝手帕旅馆", top_k=1, namespaces=["lore_test"])
    assert hits and "凯拉" in hits[0].text


def test_ingest_sqlite_rejects_non_select(kb, tmp_path):
    db = tmp_path / "x.db"
    sqlite3.connect(db).close()
    with pytest.raises(ValueError):
        ingest_sqlite(kb, db, query="DELETE FROM events")


# ── 知识范围 ──
def test_resolve_knowledge_namespaces_default_and_scoped():
    assert resolve_knowledge_namespaces(None) == DEFAULT_RETRIEVAL_NAMESPACES
    card = CharacterCard(name="x", knowledge_scope=["lore_x", "events"])
    assert resolve_knowledge_namespaces(card) == ["lore_x", "events"]


def test_resolve_knowledge_namespaces_isolates_events_by_character():
    # 无角色：公共 events（向后兼容）
    assert resolve_knowledge_namespaces(None) == DEFAULT_RETRIEVAL_NAMESPACES
    # 有角色：events 映射为 events:<cid>，web/episodic 不变
    assert resolve_knowledge_namespaces(None, "char_x") == [
        "events:char_x", "web", "episodic"
    ]
    # 角色卡知识范围中的 events 占位同样映射
    card = CharacterCard(name="x", knowledge_scope=["lore_x", "events"])
    assert resolve_knowledge_namespaces(card, "char_x") == ["lore_x", "events:char_x"]


def test_lore_namespace_sanitized():
    assert lore_namespace("wu ming/zhe") == "lore_wu_ming_zhe"


# ── RAG 上下文注入 ──
def test_build_rag_context_empty_and_lowscore():
    assert build_rag_context([]) == ""
    noise = [RetrievedChunk(text="无关", score=0.01, metadata={})]
    assert build_rag_context(noise) == ""


def test_build_roleplay_prompt_injects_tone_and_context():
    card = CharacterCard(name="无名者", personality="平静", tone="轻描淡写，多用蛾的意象")
    chunks = [
        RetrievedChunk(
            text="发条装置令大脑自动生成人格。",
            score=0.8,
            metadata={"namespace": "lore_wu_ming_zhe", "section": "设定 > 发条装置"},
        )
    ]
    prompt = build_roleplay_prompt(
        card_json=None,
        fallback_prompt=None,
        default_card=card.model_dump_json(),
        message="发条装置是什么？",
        chunks=chunks,
        card=card,
    )
    assert "无名者" in prompt
    assert "【语气要求】" in prompt and "蛾的意象" in prompt
    assert "【角色资料库" in prompt and "发条装置令大脑自动生成人格" in prompt
    assert "角色设定资料" in prompt  # 来源标签


# ── 角色卡新字段持久化 ──
def test_character_store_persists_tone_and_scope(tmp_path):
    store = CharacterStore(persist_dir=tmp_path / "chars")
    cid = store.create({"name": "测试角", "tone": "冷淡", "knowledge_scope": ["lore_a"]})
    store2 = CharacterStore(persist_dir=tmp_path / "chars")
    card = store2.get(cid)
    assert card.tone == "冷淡"
    assert card.knowledge_scope == ["lore_a"]
    # 更新其它字段不应丢失 tone/scope
    store2.update(cid, {"personality": "新性格"})
    card2 = CharacterStore(persist_dir=tmp_path / "chars").get(cid)
    assert card2.tone == "冷淡" and card2.knowledge_scope == ["lore_a"]


# ── 混合检索（稠密 + BM25 稀疏重排）──
def test_bm25_sparse_scores_reward_keyword_overlap():
    bm = _BM25([
        "维尔汀 司辰 徽章 藏在 胸口",
        "今天 天气 晴朗 适合 散步",
    ])
    s = bm.scores("维尔汀 司辰 徽章")
    assert s[0] > s[1]  # 稀疏信号：关键词命中越高分越高


def test_hybrid_rerank_promotes_exact_keyword_doc():
    """纯稠密（特征哈希，按 token 出现计数）会被高频常见词主导，
    把只含常见词干扰文档排在含专有名词的目标文档之前；
    BM25 按 IDF 加权（背景文档压低常见词权重），混合重排后目标文档回到首位。"""
    kb = KnowledgeBase(HashingEmbedder(dim=512))
    kb.add(
        [
            "维尔汀的司辰徽章是反转者的信物，藏在无名者胸口。",  # 目标：含稀有专有名词
            "解释一下含义，请说明这个东西有什么用，它有什么意义，请解释一下。",  # 干扰：常见词多
            "请解释一下含义，这个东西有什么用，它有什么意义，你说一下是什么意思。",  # 背景：压低常见词 IDF
        ],
        namespace="lore_test",
    )
    query = "维尔汀的司辰徽章有什么含义？请解释一下。"
    dense_top = kb.search(query, top_k=1, namespaces=["lore_test"], hybrid=False)[0]
    # 关键词权重调高，验证 BM25 稀疏信号能把专有名词目标顶到首位。
    # （注：特征哈希稠密分被常见词主导、严重偏离语义；生产用语义嵌入器时稠密本身
    #   已对齐关键词，默认 alpha=0.3 即可。此处用更高 alpha 显式验证 BM25 可纠正
    #   错误的稠密排序这一机制——实测 α≥0.7 时目标回到首位。）
    hybrid_top = kb.search(
        query, top_k=1, namespaces=["lore_test"], hybrid=True, hybrid_alpha=0.7
    )[0]
    # 纯稠密被常见词干扰文档带偏（目标不在首位）
    assert "司辰徽章" not in dense_top.text
    # 混合重排后命中专有名词目标
    assert "司辰徽章" in hybrid_top.text


def test_search_hybrid_flag_runs_without_error(kb):
    kb.add(["发条装置令大脑自动生成人格。", "她喜欢养蛾子。"], namespace="lore_test")
    for h in (True, False):
        hits = kb.search("发条装置", top_k=2, namespaces=["lore_test"], hybrid=h)
        assert len(hits) >= 1


# ── PDF / DOCX 导入 ──
def test_ingest_pdf_roundtrip(tmp_path):
    pdf = tmp_path / "doc.pdf"
    _write_minimal_pdf(pdf, "The pocket watch mechanism switches personalities.")
    kb = KnowledgeBase(HashingEmbedder(dim=64), persist_dir=tmp_path / "kb")
    r = ingest_file(kb, pdf, namespace="lore_pdf")
    assert r["format"] == "pdf" and r["stored"] >= 1
    hits = kb.search("pocket watch mechanism", top_k=1, namespaces=["lore_pdf"])
    assert hits and "pocket watch" in hits[0].text


def test_ingest_docx_roundtrip(tmp_path):
    from docx import Document

    doc = tmp_path / "doc.docx"
    d = Document()
    d.add_paragraph("发条装置是无名者身份切换的核心机制。")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "代号"
    t.rows[0].cells[1].text = "格蕾丝"
    d.save(str(doc))
    kb = KnowledgeBase(HashingEmbedder(dim=64), persist_dir=tmp_path / "kb")
    r = ingest_file(kb, doc, namespace="lore_docx")
    assert r["format"] == "docx" and r["stored"] >= 1
    hits = kb.search("发条装置 身份切换", top_k=1, namespaces=["lore_docx"])
    assert hits and "发条装置" in hits[0].text
